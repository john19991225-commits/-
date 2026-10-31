"""Build the formal 18-slide Smart Warehouse AI Agent proposal deck."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "presentations" / "Smart_Warehouse_AI_Agent_Proposal.pptx"

NAVY = RGBColor(15, 35, 58)
NAVY2 = RGBColor(25, 56, 82)
TEAL = RGBColor(0, 137, 142)
CYAN = RGBColor(80, 194, 201)
ORANGE = RGBColor(238, 143, 61)
RED = RGBColor(193, 65, 69)
GREEN = RGBColor(39, 154, 107)
INK = RGBColor(35, 45, 55)
MUTED = RGBColor(91, 107, 120)
LIGHT = RGBColor(242, 247, 249)
WHITE = RGBColor(255, 255, 255)
GRAY = RGBColor(211, 222, 228)


def set_run(run, size=20, bold=False, color=INK, font="Microsoft JhengHei"):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def textbox(slide, x, y, w, h, text="", size=20, bold=False, color=INK,
            align=PP_ALIGN.LEFT, fill=None, margin=0.08, valign=MSO_ANCHOR.TOP):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    set_run(r, size, bold, color)
    if fill is not None:
        shape.fill.solid(); shape.fill.fore_color.rgb = fill
        shape.line.color.rgb = fill
    return shape


def box(slide, x, y, w, h, title, body="", fill=WHITE, accent=TEAL, title_size=18,
        body_size=14, radius=True):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = fill
    s.line.color.rgb = accent; s.line.width = Pt(1.5)
    tf = s.text_frame; tf.clear(); tf.margin_left = tf.margin_right = Inches(.15)
    tf.margin_top = tf.margin_bottom = Inches(.10)
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = title; set_run(r, title_size, True, accent)
    if body:
        p2 = tf.add_paragraph(); p2.space_before = Pt(5)
        r2 = p2.add_run(); r2.text = body; set_run(r2, body_size, False, INK)
    return s


def connector(slide, x1, y1, x2, y2, color=TEAL, width=2):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color; c.line.width = Pt(width)
    c.line.end_arrowhead = True
    return c


def base_slide(prs, no, title, section, source=""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill; bg.solid(); bg.fore_color.rgb = WHITE
    top = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(.17))
    top.fill.solid(); top.fill.fore_color.rgb = TEAL; top.line.fill.background()
    textbox(slide, .55, .28, 11.7, .55, title, 25, True, NAVY)
    textbox(slide, 11.75, .32, .9, .35, f"{no:02d}", 15, True, TEAL, PP_ALIGN.RIGHT)
    textbox(slide, .58, .88, 4.0, .27, section.upper(), 9, True, MUTED)
    textbox(slide, .55, 7.08, 12.1, .22, source, 7.5, False, MUTED)
    return slide


def bullets(slide, x, y, w, h, items, size=18, color=INK, bullet_color=TEAL, gap=8):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = Inches(.08); tf.margin_right = Inches(.05)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap); p.level = 0
        r = p.add_run(); r.text = "● "; set_run(r, size-2, True, bullet_color)
        r2 = p.add_run(); r2.text = item; set_run(r2, size, False, color)
    return shape


def add_notes(slide, text):
    try:
        slide.notes_slide.notes_text_frame.text = text
    except Exception:
        pass


def add_table(slide, x, y, w, h, data, widths=None, font_size=10):
    rows, cols = len(data), len(data[0])
    table = slide.shapes.add_table(rows, cols, Inches(x), Inches(y), Inches(w), Inches(h)).table
    if widths:
        for i, width in enumerate(widths): table.columns[i].width = Inches(width)
    for r in range(rows):
        for c in range(cols):
            cell = table.cell(r, c); cell.text = str(data[r][c])
            cell.margin_left = cell.margin_right = Inches(.05)
            cell.margin_top = cell.margin_bottom = Inches(.03)
            cell.fill.solid(); cell.fill.fore_color.rgb = NAVY if r == 0 else (LIGHT if r % 2 else WHITE)
            cell.border_left = None
            for p in cell.text_frame.paragraphs:
                for run in p.runs: set_run(run, font_size if r else font_size+1, r == 0, WHITE if r == 0 else INK)
    return table


def build():
    prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)

    # 1 Cover
    s = prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb = NAVY
    s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.65), Inches(.72), Inches(.16), Inches(5.85)).fill.solid()
    bar = s.shapes[-1]; bar.fill.fore_color.rgb = TEAL; bar.line.fill.background()
    textbox(s, 1.18, 1.05, 10.9, 1.05, "Smart Warehouse AI Agent", 34, True, WHITE)
    textbox(s, 1.2, 2.05, 10.7, .75, "智慧倉儲單據辨識與庫存異常決策代理人", 26, True, CYAN)
    textbox(s, 1.2, 3.05, 10.5, .55, "Final Project Proposal", 20, True, ORANGE)
    textbox(s, 1.2, 3.72, 10.5, .55, "AI Agent  ×  Domain Knowledge  ×  Data Analytics  ×  Real-World Application", 15, False, WHITE)
    for i, (label, xx, cc) in enumerate([("DOCUMENT",1.2,TEAL),("DATA",3.45,CYAN),("AGENT",5.7,ORANGE),("HUMAN",7.95,GREEN)]):
        box(s, xx, 4.65, 1.75, .72, label, "", NAVY2, cc, 13)
        if i < 3: connector(s, xx+1.75, 5.01, xx+2.18, 5.01, WHITE, 1.5)
    textbox(s, 1.2, 6.28, 10.5, .45, "組員／課程／指導教師：請於正式繳交前填寫　｜　2026-10-31", 11, False, GRAY)
    add_notes(s, "本專題不是只做 OCR 或聊天機器人，而是由倉儲單據出發，查詢庫存、套用規則、辨識異常、保留證據，並在高風險情況交由人員確認。")

    # 2 Motivation
    s = base_slide(prs,2,"Research Motivation｜研究動機","Introduction","Source: Michel (2025), Warehouse/DC Operations Survey, n=101；產業調查，非全球母體推論。")
    for val,label,x,c in [("87%","使用 WMS",.7,TEAL),("51%","仍人工蒐集生產力資料",4.65,ORANGE),("76%","使用 inventory accuracy 指標",8.6,CYAN)]:
        box(s,x,1.45,3.25,2.0,val,label,LIGHT,c,30,15)
    textbox(s,.8,4.05,11.6,.55,"數位化 ≠ 資料可靠 ≠ 異常可解釋",26,True,NAVY,PP_ALIGN.CENTER)
    bullets(s,1.3,4.85,10.8,1.55,["前端單據仍可能含 OCR、缺欄、單位與版型錯誤。","決策需要文件、庫存、規則、證據與人工責任共同成立。"],18)
    add_notes(s,"即使倉庫使用 WMS，前端資料輸入與核對仍可能依賴人工。調查樣本只有 101，不外推全球；它只用來說明實務問題仍存在。")

    # 3 Problem
    s=base_slide(prs,3,"Problem Definition｜從文件到安全決策","Introduction","Scope: recommendation-only；不直接寫正式 WMS/ERP。")
    stages=[("INPUT","圖片／PDF\n庫存 snapshot\n主檔／歷史",.65,TEAL),("UNCERTAINTY","版型／OCR error\n缺欄／資料衝突\n過期 snapshot",3.45,ORANGE),("DOMAIN CHECK","duplicate\n數量／單位\n倉別／SKU",6.25,CYAN),("OUTPUT","欄位＋異常\n建議＋evidence\nhuman gate",9.05,GREEN)]
    for i,(t,b,x,c) in enumerate(stages):
        box(s,x,1.65,2.45,3.35,t,b,WHITE,c,17,16)
        if i<3: connector(s,x+2.45,3.32,x+2.75,3.32,c)
    textbox(s,.75,5.45,11.8,.75,"研究單位不是「辨識多少字」，而是「證據是否足以支持可追溯的倉儲建議」。",21,True,NAVY,PP_ALIGN.CENTER,LIGHT)
    add_notes(s,"輸入包含文件與外部狀態；輸出必須有結構欄位、異常、建議及證據。沒有足夠證據時，正確行為是停止或升級。")

    # 4 RQ
    s=base_slide(prs,4,"Research Questions & Measurable Objectives","Introduction","RQ 與指標需在 test 前凍結；正式實驗目前 Not Run。")
    rq=[("RQ1 文件理解","CER／Field EM／Document EM",TEAL),("RQ2 異常決策","Precision／Recall／macro-F1",ORANGE),("RQ3 Agent 價值","工具選擇／完成率／人工介入",CYAN),("RQ4 Trustworthiness","evidence coverage／false auto-accept",GREEN)]
    for i,(a,b,c) in enumerate(rq):
        y=1.35+i*1.28; box(s,.8,y,3.4,.86,a,"",LIGHT,c,18); connector(s,4.2,y+.43,4.75,y+.43,c); box(s,4.75,y,7.65,.86,b,"可重現、可稽核、學生電腦可執行",WHITE,c,17,13)
    add_notes(s,"四個問題分別回答辨識、異常、Agent 增益與可信任性，避免只有 demo 而沒有可驗證研究問題。")

    # 5 related work with article table condensed
    s=base_slide(prs,5,"Related Work｜10 篇核心研究與 Research Gap","Related Work","AI: [1]–[5]；Warehouse: [6]–[9], [13]。詳細八欄表見 README/proposal.md。")
    data=[["Ref","Year","Problem / Method","Domain / Dataset","Finding / Limitation"],
          ["[1] DocAgent","2025","multi-agent retrieval + reviewer","2 long-doc benchmarks","支援交叉核對；非倉儲"],
          ["[2] AID-Agent","2025","OCR + LLM agent + tools","44 industrial PDFs","工具化抽取；小型案例"],
          ["[3] Domain IE","2025","human/rules/LLM comparison","real-estate case","存在成本與維護權衡"],
          ["[4][5] SROIE/DocVQA","2019/21","OCR/KIE/VQA benchmarks","receipts/documents","無庫存與決策標註"],
          ["[6] IoT review","2025","systematic review","warehouse/inventory IoT","實務導入仍有限"],
          ["[7] Inventory anomaly","2025","ensemble（摘要）","12,358 SKUs（摘要）","F1 87.3%（摘要）；全文待查"],
          ["[8] AI picking","2025","case + simulation + interviews","warehouse picking","效益與導入挑戰並存"],
          ["[9][13] Smart warehouse/DT","2026","SLR / review framework","warehouse ecosystems","整合視角；非 document agent"]]
    add_table(s,.55,1.22,12.2,4.5,data,[1.4,.65,3.25,3.1,3.8],9)
    textbox(s,.65,5.95,12.0,.76,"GAP  文件 AI 多停在抽取／問答；倉儲 AI 多假設資料已結構化。\nOUR PROJECT  document → inventory context → domain rules → evidence-backed decision → human gate",16,True,NAVY,PP_ALIGN.CENTER,LIGHT)
    add_notes(s,"十篇核心研究維持 AI 與 Warehouse 各半。表中只陳述來源明示結果；跨研究數值不可直接比較。")

    # 6 mindsets
    s=base_slide(prs,6,"Five Expert Mindsets｜領域知識如何進入設計","Deep Research","Evidence: GS1 [10], NIST [12], DeHoratius & Raman [14], Destro [DR16], Chuang & Oliva [DR20].")
    minds=[("1","流程與狀態守恆","before → event → after",TEAL),("2","唯一識別與追溯","ID + provenance",CYAN),("3","Fitness for Use","品質依任務與風險",ORANGE),("4","例外與風險優先","把人工留給高風險",RED),("5","閉環校正與根因","異常 ≠ 根因",GREEN)]
    for i,(n,t,b,c) in enumerate(minds):
        x=.7+i*2.5; box(s,x,1.55,2.15,3.75,n,t+"\n\n"+b,WHITE,c,28,15)
    textbox(s,1.0,5.7,11.3,.7,"Design translation：snapshot freshness｜evidence IDs｜quality gates｜human escalation｜append-only audit",17,True,NAVY,PP_ALIGN.CENTER,LIGHT)
    add_notes(s,"五個 mindset 把 OCR 工具轉成 domain agent。例如負庫存只是 observation，不代表 Agent 已經知道 root cause。")

    # 7 controversies
    s=base_slide(prs,7,"Three Professional Controversies","Deep Research","Consensus is evidence synthesis, not a claim that every study agrees. Thresholds remain local empirical questions.")
    debates=[("高度自動化","Human-in-the-loop","Risk-based automation",TEAL),("Deterministic Rules","Learning Agent","Hybrid responsibility boundary",ORANGE),("Real operational data","Synthetic / privacy data","Mixed evidence + honest validity",CYAN)]
    for i,(a,b,c,col) in enumerate(debates):
        y=1.35+i*1.6; box(s,.7,y,3.0,.95,a,"",WHITE,col,17); connector(s,3.7,y+.48,4.4,y+.48,col); box(s,4.4,y,3.0,.95,b,"",WHITE,col,17); connector(s,7.4,y+.48,8.1,y+.48,col); box(s,8.1,y,4.5,.95,"CURRENT CONSENSUS",c,LIGHT,col,13,15)
    textbox(s,.75,6.25,11.85,.48,"Unresolved：threshold、人工負荷、domain shift、distribution gap、成本效益。",16,True,RED,PP_ALIGN.CENTER)
    add_notes(s,"不是二選一。共識是依風險採 hybrid 設計；最佳門檻、人工負荷與 domain shift 仍要由本地資料回答。")

    # 8 questions
    s=base_slide(prs,8,"Ten Deep Questions｜以失敗情境檢查理解","Deep Research","完整分析、證據與設計啟示見 deep_research.md Q3。")
    qs=["高 OCR accuracy 為何仍會錯決策？","文件與 snapshot 時間不同怎麼辦？","自動通過與危險錯放如何權衡？","OCR、barcode、主檔衝突時信誰？","什麼是充分 evidence？","為何保留 OCR + Rules baseline？","為何 planner 必須 bounded？","人工核准後發現錯誤怎麼追？","新版型低信心時能降 threshold 嗎？","盤點差異為何不能直接定因？"]
    for i,q in enumerate(qs):
        x=.65+(i%2)*6.15; y=1.25+(i//2)*1.05
        box(s,x,y,5.75,.78,f"Q{i+1:02d}",q,WHITE,[TEAL,CYAN,ORANGE,GREEN,RED][i%5],13,13)
    textbox(s,.8,6.62,11.7,.3,"核心原則：觀測、推論、未知與人工處置必須分開保存。",15,True,NAVY,PP_ALIGN.CENTER)
    add_notes(s,"報告時可深入第 1、3、4、7 題。十題涵蓋 Why、What-if、Trade-off、Evidence、Comparison 與 Failure Case。")

    # 9 agent
    s=base_slide(prs,9,"Proposed Domain AI Agent｜能力與責任邊界","Proposed Scheme","Agent ≠ fixed pipeline；只有能依 state 選工具、補查、停止或升級者才稱 Agent。")
    comps=[("PLANNER","task → plan\nallowlist / max steps",TEAL),("TOOLS","OCR / parser / SQL\nanalytics / rules",CYAN),("REASONING","facts → conflicts\nevidence-linked rationale",ORANGE),("DECISION","severity × confidence\ncoverage × freshness",RED),("HUMAN","approve / edit / reject\nreprocess",GREEN)]
    for i,(t,b,c) in enumerate(comps):
        x=.45+i*2.56; box(s,x,1.55,2.25,3.35,t,b,WHITE,c,15,14)
        if i<4: connector(s,x+2.25,3.22,x+2.52,3.22,c)
    bullets(s,.8,5.2,11.7,1.25,["LLM：任務映射與受控工具選擇；不得創造數量、SKU 或政策。","Rules／SQL／statistics：可重現計算；最終安全 gate 與正式寫入權限不交給 LLM。"],15)
    add_notes(s,"OCR、SQL 和規則仍是普通可測試工具。Planner 的價值在於依中間結果選擇下一步或安全停止。")

    # 10 architecture, native diagram
    s=base_slide(prs,10,"System Architecture｜①–⑧ 可測試模組","Proposed Scheme","Target design: local-first, read-only inventory, evidence-first, human-gated, append-only audit.")
    mods=[("① User","scenario / role",TEAL),("② Input","task + image/PDF",TEAL),("③ Planner","bounded tool plan",ORANGE),("④ Retrieval","OCR / parse / SQL",CYAN),("⑤ Analytics","duplicate / reconcile",CYAN),("⑥ Knowledge","rules / master / GS1",ORANGE),("⑦ Decision","risk + evidence gate",RED),("⑧ Interaction","result / review / audit",GREEN)]
    coords=[]
    for i,(t,b,c) in enumerate(mods):
        row=i//4; col=i%4; x=.45+col*3.18; y=1.35+row*2.3
        coords.append((x,y)); box(s,x,y,2.75,1.48,t,b,WHITE,c,16,13)
    for i in range(3): connector(s,coords[i][0]+2.75,coords[i][1]+.74,coords[i+1][0],coords[i+1][1]+.74,mods[i][2])
    connector(s,coords[3][0]+1.37,coords[3][1]+1.48,coords[7][0]+1.37,coords[7][1],CYAN)
    for i in range(7,4,-1): connector(s,coords[i][0],coords[i][1]+.74,coords[i-1][0]+2.75,coords[i-1][1]+.74,mods[i][2])
    connector(s,coords[4][0]+1.37,coords[4][1],coords[0][0]+1.37,coords[0][1]+1.48,MUTED,1)
    textbox(s,.75,6.25,11.85,.45,"Human reprocess loop returns to Planner with the same root trace; no formal WMS write in MVP.",14,True,NAVY,PP_ALIGN.CENTER,LIGHT)
    add_notes(s,"架構把 perception、analytics、rules、decision 與 interaction 分開，使錯誤可定位，每個模組可單獨驗證。")

    # 11 flowchart
    s=base_slide(prs,11,"Flowchart｜Normal Path, Exception, Human Loop","Proposed Scheme","Normal=teal；Decision=orange；Exception=red；Human interaction=green.")
    normal=[("INPUT",.4),("PROCESS",1.75),("OCR",3.1),("VALIDATE",4.45),("ANALYTICS",5.8),("DOMAIN",7.15),("DECISION",8.5),("OUTPUT",9.85)]
    for i,(t,x) in enumerate(normal):
        c=ORANGE if t in ("VALIDATE","DECISION") else (GREEN if t=="OUTPUT" else TEAL)
        box(s,x,1.3,1.12,.72,t,"",WHITE,c,11)
        if i<7: connector(s,x+1.12,1.66,x+1.31,1.66,c,1.5)
    exceptions=[("OCR fail → retry once",2.75,2.55), ("missing critical field",4.3,3.45), ("duplicate / mismatch",6.0,2.55), ("stale source / rule error",7.7,3.45)]
    for t,x,y in exceptions: box(s,x,y,2.25,.72,"EXCEPTION",t,WHITE,RED,11,11)
    box(s,5.1,5.0,3.15,1.0,"HUMAN REVIEW","approve / edit / reject / reprocess",LIGHT,GREEN,14,12)
    for _,x,y in exceptions: connector(s,x+1.1,y+.72,6.68,5.0,RED,1)
    connector(s,5.1,5.5,1.0,1.3,GREEN,1.5)
    textbox(s,.7,6.35,11.9,.46,"Fail-safe：來源不可用 ≠ 查無資料；模糊候選 ≠ confirmed duplicate；reprocess 不覆寫舊 trace。",14,True,NAVY,PP_ALIGN.CENTER,LIGHT)
    add_notes(s,"失敗不是空值成功。OCR retry 只做一次，缺欄、資料失效、規則衝突、重複與庫存不一致都進人工。")

    # 12 algorithm
    s=base_slide(prs,12,"Algorithm 1｜Bounded Planning + Evidence Gate","Proposed Scheme","Complete 92-line pseudocode: algorithm.md；slide shows critical control points.")
    code=("01  validate input; start append-only trace\n"
          "04  task ← UNDERSTAND_TASK(q, allowed scenarios)\n"
          "09  plan ← BUILD_BOUNDED_PLAN(allowed tools, max steps)\n"
          "15  OCR → retry once → ESCALATE if unusable\n"
          "24  parse typed schema; link fields to OCR spans\n"
          "33  read-only inventory / history / master queries\n"
          "43  exact duplicate ≠ fuzzy candidates\n"
          "50  evaluate versioned warehouse rules\n"
          "57  reason from evidence only\n"
          "62  confidence + coverage + risk\n"
          "67  duplicate / mismatch / uncertainty → HUMAN REVIEW\n"
          "79  ASSERT_NO_WMS_OR_INVENTORY_WRITE_OCCURRED\n"
          "91  append trace and return")
    textbox(s,.65,1.25,7.35,5.45,code,14,False,WHITE,PP_ALIGN.LEFT,NAVY2,.18)
    box(s,8.35,1.45,4.2,1.2,"INVARIANT 1","All tools are allowlisted and schema-valid.",WHITE,TEAL,15,13)
    box(s,8.35,3.0,4.2,1.2,"INVARIANT 2","Every material claim has evidence ID.",WHITE,ORANGE,15,13)
    box(s,8.35,4.55,4.2,1.2,"INVARIANT 3","High risk / uncertainty cannot auto-write.",WHITE,RED,15,13)
    add_notes(s,"演算法重點不是 92 行本身，而是每個 exit condition 都 fail-safe，並且原型沒有自動寫正式庫存的能力。")

    # 13 MVP
    s=base_slide(prs,13,"MVP Implementation｜What Exists vs. What Is Planned","Technical Feasibility","Automated tests: 6/6 synthetic fixtures passed; this is not formal experiment evidence.")
    cols=[("IMPLEMENTED",["Streamlit upload","mock/sidecar OCR","typed parser","SQLite seed/query","exact duplicate + rules","human confirm + audit"],GREEN),
          ("PARTIAL",["bounded planner","embedded-text PDF","confidence gate","target architecture mapping"],ORANGE),
          ("PLANNED",["real/scanned OCR benchmark","LLM adaptive planner","fuzzy duplicate","Google Sheets adapter","formal WMS integration"],RED)]
    for i,(t,items,c) in enumerate(cols):
        x=.55+i*4.25; box(s,x,1.25,3.8,.72,t,"",LIGHT,c,17)
        bullets(s, x+.12, 2.15, 3.55, 3.95, items, size=15, bullet_color=c, gap=6)
    textbox(s,.7,6.25,11.9,.48,"Python 3.11 + Streamlit + SQLite + Pandas｜一般學生電腦，不訓練大型模型。",16,True,NAVY,PP_ALIGN.CENTER,LIGHT)
    add_notes(s,"這頁誠實區分程式存在、部分完成與規劃。六個測試只支持合成 fixture 的端到端行為。")

    # 14 dataset
    s=base_slide(prs,14,"Dataset Design｜No Leakage, No Sensitive Documents","Experiments","Main dataset target is planned, not collected: 300–500 documents, ≥3 types, ≥5 templates.")
    layers=[("PUBLIC PRETEST","SROIE / DocVQA\npipeline sanity only",TEAL,1.1),("MAIN DATASET","anonymous / synthetic\nChinese + English warehouse docs",ORANGE,2.7),("PAIRED CONTEXT","immutable snapshots\nmaster + history + gold decision",GREEN,4.3)]
    for t,b,c,y in layers: box(s,1.1,y,5.2,1.15,t,b,WHITE,c,16,14)
    box(s,7.1,1.25,5.1,3.85,"GROUP SPLIT","vendor_id + template_id\n\n60 / 20 / 20\nseed = 20261024\n\nSame original + augmentation\nMUST stay in one split",LIGHT,NAVY,17,18)
    textbox(s,7.15,5.45,5.0,.72,"Actual sample size / license / class distribution：NOT RUN／待查證",14,True,RED,PP_ALIGN.CENTER)
    add_notes(s,"公開收據只驗證 pipeline，不代表臺灣倉儲。主資料尚未取得，不能填入假樣本數或結果。")

    # 15 baselines
    s=base_slide(prs,15,"Three Baselines｜Fair Paired Comparison","Experiments","All baselines use the same test documents and immutable database snapshots.")
    bases=[("B0 MANUAL","人工輸入＋人工核對","Operational reference",TEAL),("B1 FIXED","OCR＋Fixed Rules","Non-agent control",ORANGE),("B2 AGENT","OCR＋Domain AI Agent","Proposed method",GREEN)]
    for i,(a,b,c,col) in enumerate(bases):
        x=.65+i*4.2; box(s,x,1.35,3.75,2.25,a,b+"\n\n"+c,WHITE,col,20,16)
    textbox(s,.8,4.15,11.7,.55,"COMMON：OCR｜parser｜rules｜thresholds｜hardware｜test manifest",18,True,NAVY,PP_ALIGN.CENTER,LIGHT)
    bullets(s,1.1,5.0,11.2,1.25,["B1 固定順序、不 replan；B2 才能依中間 state 選工具或升級。","若 B1/B2 使用不同 OCR 或規則，就無法把差異歸因於 Agent planning。"],15)
    add_notes(s,"B1 是關鍵控制組。三者必須用相同資料；B1/B2 還要共用 OCR、parser、rules 與 threshold。")

    # 16 metrics
    s=base_slide(prs,16,"Evaluation Metrics & Procedure","Experiments","Primary endpoint proposed: anomaly macro-F1（teacher confirmation required）；all results Not Run.")
    metrics=[("DOCUMENT","CER / Character Accuracy\nField EM / Document EM",TEAL),("ANOMALY","Precision / Recall / macro-F1\nDuplicate Accuracy + F1",ORANGE),("OPERATIONS","Mean / median / P95 time\nHuman intervention",CYAN),("AGENT","Tool precision / recall\nSafe tool / task completion",GREEN)]
    for i,(t,b,c) in enumerate(metrics):
        x=.55+i*3.15; box(s,x,1.25,2.85,1.75,t,b,WHITE,c,15,13)
    steps=["1 Annotate + adjudicate","2 Freeze split/config","3 Run paired B0/B1/B2","4 Bootstrap 95% CI","5 Error analysis"]
    for i,st in enumerate(steps):
        x=.65+i*2.5; box(s,x,3.65,2.15,.9,st,"",LIGHT,NAVY,13)
        if i<4: connector(s,x+2.15,4.1,x+2.42,4.1,NAVY,1)
    add_table(s,1.15,5.1,11.0,1.05,[["Metric","B0","B1","B2","Status"],["macro-F1 / time / intervention / completion","—","—","—","NOT RUN"]],[5.9,1.1,1.1,1.1,1.8],11)
    add_notes(s,"所有方案按同一 manifest 配對比較。Agent 成功還要求選對工具、沒有越權、證據完整且 disposition 正確。")

    # 17 contributions
    s=base_slide(prs,17,"Expected Contributions & Feasibility","Conclusion","Expected outcomes are hypotheses to test, not observed findings.")
    cont=[("C1","Research framing","OCR → domain decision agent",TEAL),("C2","System design","evidence-first + human-gated",ORANGE),("C3","Evaluation","fair B0/B1/B2 + safety metrics",CYAN),("C4","Artifact","student-computer open-source MVP",GREEN)]
    for i,(n,t,b,c) in enumerate(cont):
        x=.55+i*3.15; box(s,x,1.35,2.85,2.2,n,t+"\n\n"+b,WHITE,c,24,14)
    textbox(s,.8,4.15,11.7,.65,"EXPECTED：B2 可能改善跨來源異常處理與 auditability；代價可能是 latency、complexity 與人工 threshold 調整。",17,True,NAVY,PP_ALIGN.CENTER,LIGHT)
    textbox(s,.8,5.25,11.7,.65,"NOT CLAIMED：尚無正式 accuracy、F1、time saving、ROI 或「優於所有方法」的結果。",17,True,RED,PP_ALIGN.CENTER)
    add_notes(s,"預期貢獻必須能被反駁、量測，並在失敗時安全停止。是否改善 F1 或時間只能由後續實驗回答。")

    # 18 refs and roadmap
    s=base_slide(prs,18,"References, Submission Status & Next Steps","Conclusion","Repository: https://github.com/john19991225-commits/-")
    refs=("[1] Sun et al. (2025), DocAgent, EMNLP. doi:10.18653/v1/2025.emnlp-main.893\n"
          "[2] Li, Conen & Aller (2025), AID-Agent, REALM. doi:10.18653/v1/2025.realm-1.6\n"
          "[3] Principe et al. (2025), LDK, ACL Anthology 2025.ldk-1.25\n"
          "[6] Puga Mendoza et al. (2025), Procedia CS. doi:10.1016/j.procs.2025.01.185\n"
          "[8] Fakhrai Rad et al. (2025), Review of Managerial Science. doi:10.1007/s11846-025-00858-1\n"
          "[9] El Kadmiri et al. (2026), Cleaner Logistics and Supply Chain. doi:10.1016/j.clscn.2026.100340\n"
          "[10] GS1 Logistic Label Guideline 1.3 (2019)｜[12] NIST AI RMF 1.0 (2023)\n"
          "[13] Alam et al. (2026), ICT Express. doi:10.1016/j.icte.2026.01.009\n"
          "Full verified list: README.md / proposal.md / deep_research.md")
    textbox(s,.55,1.2,7.5,5.55,refs,12.5,False,INK,PP_ALIGN.LEFT,LIGHT,.18)
    roadmap=[("NOW","Proposal + Design + MVP",GREEN),("NEXT","teacher confirms protocol",TEAL),("PILOT","30–50 compliant cases",ORANGE),("FREEZE","manifest + versions",NAVY),("RUN","B0 / B1 / B2",RED)]
    for i,(a,b,c) in enumerate(roadmap):
        y=1.25+i*1.08; box(s,8.35,y,4.2,.78,a,b,WHITE,c,13,12)
    add_notes(s,"目前完成 Proposal、Deep Research、Design、PPTX 與合成資料 MVP。下一個 gate 是教師確認、合規資料與 pilot，不是先填漂亮數字。")

    # Core properties and save
    prs.core_properties.title = "Smart Warehouse AI Agent — Final Project Proposal"
    prs.core_properties.subject = "Domain AI Agent, document understanding, inventory anomaly decision"
    prs.core_properties.author = "Smart Warehouse AI Agent Project Team"
    prs.core_properties.comments = "Formal proposal deck. Experimental results are Not Run / Expected Outcomes."
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f"Created: {OUT}")
    print(f"Slides: {len(prs.slides)}")


if __name__ == "__main__":
    build()
